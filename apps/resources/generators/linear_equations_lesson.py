"""Builds the "Белгисиз бир жагында: 1-кадам" (one-step equations, variable
on one side) teaching presentation, following Singapore Math's
Concrete -> Pictorial -> Abstract (CPA) sequence: every new equation type
is shown first as a balance scale or bar model, then as the matching
symbolic steps, so the abstract algebra is always anchored to something
the student can see. Content and layout reuse the same palette, font and
slide size as pptx_builder.py so a downloaded lesson deck looks like it
belongs to the same site as the recall packs.

This is a fixed, hand-authored lesson (not a parameterized/randomized
generator like recall_pack.py) - there is one build_lesson() entry point
and no seed argument.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from .pptx_builder import (
    SLIDE_W, SLIDE_H, MARGIN, NAVY, BLUE, GRAY, LIGHT_GRAY, FONT,
    _add_textbox,
)

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAN_FILL = RGBColor(0xDC, 0xE6, 0xEB)
UNKNOWN_FILL = BLUE
KNOWN_FILL = RGBColor(0x94, 0xA3, 0xB8)

TITLE_H = Emu(900000)


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _header(slide, text):
    _add_textbox(slide, MARGIN, Emu(260000), SLIDE_W - 2 * MARGIN, TITLE_H,
                 text, size=26, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


def _title_slide(prs, title_lines, subtitle):
    slide = _blank_slide(prs)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    bg.shadow.inherit = False
    _add_textbox(slide, Emu(600000), Emu(2200000), SLIDE_W - Emu(1200000), Emu(1800000),
                 title_lines, size=34, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    _add_textbox(slide, Emu(600000), Emu(3900000), SLIDE_W - Emu(1200000), Emu(500000),
                 subtitle, size=16, bold=False, color=RGBColor(0xBF, 0xD6, 0xE8),
                 align=PP_ALIGN.CENTER)
    _add_textbox(slide, Emu(600000), SLIDE_H - Emu(700000), SLIDE_W - Emu(1200000), Emu(400000),
                 "MathKGZ", size=14, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    return slide


def _bullets(slide, left, top, width, height, lines, size=16):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = f"•  {line}"
        run.font.size = Pt(size)
        run.font.name = FONT
        run.font.color.rgb = NAVY


def _pan(slide, cx, top, width, block_labels):
    """One balance pan: a shallow trapezoid holding a row of small
    labeled squares (the "unit blocks" Singapore Math models a quantity
    or the unknown with)."""
    pan_h = Emu(260000)
    pan_top = top + Emu(260000)
    left = Emu(int(cx - width / 2))
    pts = [
        (left, pan_top), (left + width, pan_top),
        (left + width - Emu(80000), pan_top + pan_h),
        (left + Emu(80000), pan_top + pan_h),
    ]
    fb = slide.shapes.build_freeform(start_x=pts[0][0], start_y=pts[0][1])
    fb.add_line_segments(pts[1:] + [pts[0]], close=True)
    pan = fb.convert_to_shape()
    pan.fill.solid()
    pan.fill.fore_color.rgb = WHITE
    pan.line.color.rgb = LIGHT_GRAY
    pan.line.width = Pt(2)
    pan.shadow.inherit = False

    block = Emu(220000)
    gap = Emu(30000)
    n = len(block_labels)
    row_w = n * block + (n - 1) * gap
    bx = Emu(int(cx - row_w / 2))
    by = top - Emu(20000)
    for label, is_unknown in block_labels:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, block, block)
        shape.adjustments[0] = 0.15
        shape.fill.solid()
        shape.fill.fore_color.rgb = UNKNOWN_FILL if is_unknown else KNOWN_FILL
        shape.line.color.rgb = NAVY
        shape.line.width = Pt(1)
        shape.shadow.inherit = False
        tf = shape.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.name = FONT
        run.font.color.rgb = WHITE
        bx = Emu(int(bx + block + gap))


def _draw_balance(slide, left, top, width, height, left_blocks, right_blocks, caption):
    """A balance scale: a fulcrum, a level beam, and two pans, each
    holding labeled unit blocks. left_blocks/right_blocks are lists of
    (label, is_unknown) tuples, e.g. [("x", True), ("3", False)]."""
    cx = Emu(int(left + width / 2))
    fulcrum_top = top + Emu(900000)
    fulcrum_h = Emu(500000)
    fulcrum_w = Emu(260000)
    fulcrum_pts = [
        (Emu(int(cx - fulcrum_w / 2)), fulcrum_top + fulcrum_h),
        (Emu(int(cx + fulcrum_w / 2)), fulcrum_top + fulcrum_h),
        (cx, fulcrum_top),
    ]
    fb = slide.shapes.build_freeform(start_x=fulcrum_pts[0][0], start_y=fulcrum_pts[0][1])
    fb.add_line_segments(fulcrum_pts[1:] + [fulcrum_pts[0]], close=True)
    fulcrum = fb.convert_to_shape()
    fulcrum.fill.solid()
    fulcrum.fill.fore_color.rgb = NAVY
    fulcrum.line.fill.background()
    fulcrum.shadow.inherit = False

    beam_w = Emu(int(width - Emu(400000)))
    beam = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(cx - beam_w / 2)), fulcrum_top - Emu(40000),
        beam_w, Emu(40000),
    )
    beam.adjustments[0] = 0.5
    beam.fill.solid()
    beam.fill.fore_color.rgb = NAVY
    beam.line.fill.background()
    beam.shadow.inherit = False

    pan_cx_left = Emu(int(cx - beam_w / 2 + Emu(300000)))
    pan_cx_right = Emu(int(cx + beam_w / 2 - Emu(300000)))
    pan_top_y = fulcrum_top - Emu(360000)
    _pan(slide, pan_cx_left, pan_top_y, Emu(1100000), left_blocks)
    _pan(slide, pan_cx_right, pan_top_y, Emu(1100000), right_blocks)

    _add_textbox(slide, left, top + height - Emu(420000), width, Emu(400000),
                 caption, size=14, color=GRAY, align=PP_ALIGN.CENTER)


def _draw_bar_model(slide, left, top, width, height, segment_labels, whole_label, caption):
    """A single bar split into equal, labeled segments, with a brace-like
    bracket under the whole bar naming the total. Used for multiplication
    (several equal parts making a known whole) and division (a whole
    split into equal known parts)."""
    bar_w = width - Emu(400000)
    bar_h = Emu(700000)
    bar_left = Emu(int(left + (width - bar_w) / 2))
    bar_top = top + Emu(300000)
    n = len(segment_labels)
    seg_w = Emu(int(bar_w / n))

    for i, (label, is_unknown) in enumerate(segment_labels):
        seg_left = Emu(int(bar_left + i * seg_w))
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, seg_left, bar_top, seg_w, bar_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = UNKNOWN_FILL if is_unknown else KNOWN_FILL
        shape.line.color.rgb = WHITE
        shape.line.width = Pt(2)
        shape.shadow.inherit = False
        tf = shape.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.name = FONT
        run.font.color.rgb = WHITE

    outline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, bar_left, bar_top, bar_w, bar_h)
    outline.fill.background()
    outline.line.color.rgb = NAVY
    outline.line.width = Pt(2.5)
    outline.shadow.inherit = False

    brace_top = bar_top + bar_h + Emu(60000)
    _add_textbox(slide, bar_left, brace_top, bar_w, Emu(400000),
                 f"бардыгы = {whole_label}", size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    _add_textbox(slide, left, top + height - Emu(380000), width, Emu(360000),
                 caption, size=14, color=GRAY, align=PP_ALIGN.CENTER)


def _steps(slide, left, top, width, height, steps):
    """steps: list of (equation_text, annotation_text_or_None), rendered
    as a vertical stack - the annotation makes each algebraic move
    traceable back to "what did we just do to the balance/bar model"."""
    row_h = Emu(int(height / len(steps)))
    y = top
    for equation, note in steps:
        _add_textbox(slide, left, y, width, Emu(int(row_h * 0.62)),
                     equation, size=22, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        if note:
            _add_textbox(slide, left, y + Emu(int(row_h * 0.60)), width, Emu(int(row_h * 0.38)),
                         note, size=12, color=BLUE, align=PP_ALIGN.CENTER)
        y = Emu(int(y + row_h))


def _two_column(slide, left_builder, right_builder):
    col_w = Emu(int((SLIDE_W - 2 * MARGIN - Emu(200000)) / 2))
    left_x = MARGIN
    right_x = MARGIN + col_w + Emu(200000)
    top = Emu(1300000)
    height = SLIDE_H - top - Emu(500000)
    left_builder(left_x, top, col_w, height)
    right_builder(right_x, top, col_w, height)


def _closing_slide(prs):
    slide = _blank_slide(prs)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    bg.shadow.inherit = False
    _add_textbox(slide, Emu(600000), Emu(2000000), SLIDE_W - Emu(1200000), Emu(400000),
                 "Эстен чыгарба!", size=26, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    _add_textbox(
        slide, Emu(600000), Emu(2600000), SLIDE_W - Emu(1200000), Emu(2000000),
        [
            "Теңдеме — тараза: эки жагы дайым тең.",
            "Белгисизди (x) жалгыз калтыруу үчүн,",
            "экинчи жагындагы амалдын карама-каршысын",
            "эки жагына тең аткар.",
        ],
        size=18, bold=False, color=WHITE, align=PP_ALIGN.CENTER,
    )
    return slide


def build_lesson():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1. Title
    _title_slide(
        prs,
        ["Белгисиз бир жагында турган", "теңдемелер"],
        "1-кадамдуу теңдемелер  ·  7–8-класс",
    )

    # 2. Learning goal
    slide = _blank_slide(prs)
    _header(slide, "Бул сабакта эмнени үйрөнөсүң?")
    _bullets(slide, MARGIN, Emu(1300000), SLIDE_W - 2 * MARGIN, Emu(2600000), [
        "x + a = b,  x − a = b,  a·x = b,  x : a = b түрүндөгү теңдемелерди чечүү",
        "Теңдеме — таразага окшош: эки жагы ар дайым тең салмакта болушу керек",
        "Белгисизди (x) жалгыз калтыруу үчүн, эки жагынан тең бир эле амалды аткаруу",
    ], size=18)

    # 3. Concrete/Pictorial: x + 3 = 8 balanced
    slide = _blank_slide(prs)
    _header(slide, "Теңдеме — тараза")
    _draw_balance(
        slide, MARGIN, Emu(1200000), SLIDE_W - 2 * MARGIN, Emu(3600000),
        left_blocks=[("x", True), ("3", False)],
        right_blocks=[("8", False)],
        caption="Эки жак тең салмакта — демек, алар бирдей: x + 3 = 8",
    )

    # 4. Pictorial: remove 3 from both sides
    slide = _blank_slide(prs)
    _header(slide, "Тең салмакты сактоо үчүн эки жагынан бирдей санды кемитебиз")
    _draw_balance(
        slide, MARGIN, Emu(1200000), SLIDE_W - 2 * MARGIN, Emu(3600000),
        left_blocks=[("x", True)],
        right_blocks=[("5", False)],
        caption="Эки жагынан тең 3төн алып салсак, тең салмак сакталат: x = 5",
    )

    # 5. Abstract: symbolic steps for example 1
    slide = _blank_slide(prs)
    _header(slide, "Эми ушул эле кадамдарды сандар менен жазабыз")
    _steps(slide, MARGIN, Emu(1400000), SLIDE_W - 2 * MARGIN, Emu(3200000), [
        ("x + 3 = 8", None),
        ("x + 3 − 3 = 8 − 3", "эки жагынан тең 3төн кемитебиз"),
        ("x = 5", None),
    ])

    # 6. Worked example 2: subtraction type, two-column (balance + steps)
    slide = _blank_slide(prs)
    _header(slide, "2-мисал:  x − 4 = 9")

    def ex2_left(l, t, w, h):
        _draw_balance(
            slide, l, t, w, h,
            left_blocks=[("x", True)],
            right_blocks=[("9", False), ("4", False)],
            caption="x − 4 = 9  (тараза эки жагынан 4 кем турат)",
        )

    def ex2_right(l, t, w, h):
        _steps(slide, l, t, w, h, [
            ("x − 4 = 9", None),
            ("x − 4 + 4 = 9 + 4", "эки жагына тең 4төн кошобуз"),
            ("x = 13", None),
        ])

    _two_column(slide, ex2_left, ex2_right)

    # 7. Worked example 3: multiplication type, bar model
    slide = _blank_slide(prs)
    _header(slide, "3-мисал:  3x = 12")

    def ex3_left(l, t, w, h):
        _draw_bar_model(
            slide, l, t, w, h,
            segment_labels=[("x", True), ("x", True), ("x", True)],
            whole_label="12",
            caption="3 барабар бөлүк кошулса, 12 болот",
        )

    def ex3_right(l, t, w, h):
        _steps(slide, l, t, w, h, [
            ("3x = 12", None),
            ("3x ÷ 3 = 12 ÷ 3", "эки жагын тең 3кө бөлөбүз"),
            ("x = 4", None),
        ])

    _two_column(slide, ex3_left, ex3_right)

    # 8. Worked example 4: division type, bar model
    slide = _blank_slide(prs)
    _header(slide, "4-мисал:  x : 2 = 6")

    def ex4_left(l, t, w, h):
        _draw_bar_model(
            slide, l, t, w, h,
            segment_labels=[("6", False), ("6", False)],
            whole_label="x",
            caption="x экиге бөлүнгөндө, ар бир бөлүк 6",
        )

    def ex4_right(l, t, w, h):
        _steps(slide, l, t, w, h, [
            ("x : 2 = 6", None),
            ("x : 2 × 2 = 6 × 2", "эки жагын тең 2ге көбөйтөбүз"),
            ("x = 12", None),
        ])

    _two_column(slide, ex4_left, ex4_right)

    # 9. Generalization
    slide = _blank_slide(prs)
    _header(slide, "Эрежени тапкыла")
    _bullets(slide, MARGIN, Emu(1300000), SLIDE_W - 2 * MARGIN, Emu(2200000), [
        "+ кошулса  →  эки жагынан кемитебиз",
        "− алынса  →  эки жагына кошобуз",
        "· көбөйтүлсө  →  эки жагын бөлөбүз",
        ": бөлүнсө  →  эки жагын көбөйтөбүз",
    ], size=19)
    _add_textbox(
        slide, MARGIN, Emu(3600000), SLIDE_W - 2 * MARGIN, Emu(900000),
        "Эреже: белгисизди жалгыз калтыруу үчүн, экинчи жагындагы амалдын "
        "карама-каршысын эки жагына тең аткар.",
        size=16, bold=True, color=BLUE, align=PP_ALIGN.CENTER,
    )

    # 10. Guided practice
    slide = _blank_slide(prs)
    _header(slide, "Көнүгүү — өзүңүз чечиңиз")
    _bullets(slide, MARGIN, Emu(1300000), SLIDE_W - 2 * MARGIN, Emu(3200000), [
        "x + 7 = 15",
        "x − 9 = 6",
        "5x = 35",
        "x : 4 = 9",
        "x + 12 = 20",
        "2x = 18",
    ], size=20)

    # 11. Answer key
    slide = _blank_slide(prs)
    _header(slide, "Жооптор")
    _bullets(slide, MARGIN, Emu(1300000), SLIDE_W - 2 * MARGIN, Emu(3200000), [
        "x + 7 = 15  →  x = 8",
        "x − 9 = 6  →  x = 15",
        "5x = 35  →  x = 7",
        "x : 4 = 9  →  x = 36",
        "x + 12 = 20  →  x = 8",
        "2x = 18  →  x = 9",
    ], size=20)

    # 12. Summary
    _closing_slide(prs)

    return prs
