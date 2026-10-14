"""Builds a native, fully-editable PowerPoint Tarsia jigsaw: real triangle
shapes (not a raster image) plus rotated textboxes for each question/
answer, using the same geometry as tarsia_svg.py / tarsia_geom.py. python-
pptx rotation is clockwise degrees about the shape's own center, which is
exactly what text_rotation() already returns given our y-down coordinate
system - no sign flip needed.
"""
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from .tarsia_geom import (
    SQRT3,
    internal_edges,
    text_rotation,
    triangle_centroid,
    triangle_vertices,
)

INK = RGBColor(0x1F, 0x2A, 0x44)
MUTED = RGBColor(0x5B, 0x64, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xF2, 0xA3, 0x3F)
FONT = 'Calibri'

SLIDE_W = Emu(12192000)  # 13.333in, LAYOUT_WIDE
SLIDE_H = Emu(6858000)   # 7.5in


def _title_slide(prs, title, usage_lines):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(540000), Emu(450000), Emu(11100000), Emu(1000000))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.name = FONT
    p.font.color.rgb = INK

    box2 = slide.shapes.add_textbox(Emu(540000), Emu(1620000), Emu(11100000), Emu(2700000))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    for i, line in enumerate(usage_lines):
        p2 = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p2.text = line
        p2.font.size = Pt(15)
        p2.font.name = FONT
        p2.font.color.rgb = MUTED
    return slide


def _closing_slide(prs, closing_lines, closing_email):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = INK
    box = slide.shapes.add_textbox(Emu(720000), Emu(1800000), Emu(10550000), Emu(1450000))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(closing_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.name = FONT
        p.font.color.rgb = WHITE
    if closing_email:
        box2 = slide.shapes.add_textbox(Emu(720000), Emu(3420000), Emu(10550000), Emu(540000))
        tf2 = box2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = closing_email
        p2.font.size = Pt(14)
        p2.font.name = FONT
        p2.font.color.rgb = RGBColor(0xC7, 0xD2, 0xE8)


def _puzzle_slide(prs, n, pairs, s=1.8, font_size=10, inset_frac=0.13):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tris = triangle_vertices(n, s)
    edges = internal_edges(n, s)

    grid_w = n * s
    grid_h = n * s * SQRT3 / 2
    target_h = Emu(5900000)
    px_per_unit = int(target_h / grid_h)
    target_w = Emu(int(grid_w * px_per_unit))
    off_x = (SLIDE_W - target_w) // 2
    off_y = Emu(620000)

    def tx(x):
        return Emu(int(off_x + (x + grid_w / 2) * px_per_unit))

    def ty(y):
        return Emu(int(off_y + y * px_per_unit))

    for verts in tris.values():
        builder = slide.shapes.build_freeform(tx(verts[0][0]), ty(verts[0][1]))
        builder.add_line_segments([(tx(verts[1][0]), ty(verts[1][1])), (tx(verts[2][0]), ty(verts[2][1])), (tx(verts[0][0]), ty(verts[0][1]))], close=True)
        shape = builder.convert_to_shape()
        shape.fill.background()
        shape.line.color.rgb = INK
        shape.line.width = Pt(1.5)
        shape.shadow.inherit = False

    inset = inset_frac * s
    box_size_data = s * 0.42
    box_px = int(box_size_data * px_per_unit)
    for (ka, kb, mid, angle), pair in zip(edges, pairs):
        q, a = (list(pair) + ['', ''])[:2]
        for key, text in ((ka, q), (kb, a)):
            if not text:
                continue
            verts = tris[key]
            rot = text_rotation(mid, verts, angle)
            cx, cy = triangle_centroid(verts)
            mx, my = mid
            dx, dy = cx - mx, cy - my
            dist = (dx ** 2 + dy ** 2) ** 0.5 or 1
            ux, uy = dx / dist, dy / dist
            cxp = int(tx(mx + ux * inset))
            cyp = int(ty(my + uy * inset))
            left = Emu(cxp - box_px // 2)
            top = Emu(cyp - box_px // 3)
            box = slide.shapes.add_textbox(left, top, Emu(box_px), Emu(int(box_px * 2 / 3)))
            box.rotation = rot
            tf = box.text_frame
            tf.word_wrap = False
            tf.margin_left = 0
            tf.margin_right = 0
            tf.margin_top = 0
            tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = str(text)
            run.font.size = Pt(font_size)
            run.font.name = FONT
            run.font.color.rgb = INK
    return slide


def build_tarsia_pptx(n, pairs, title, usage_lines,
                       closing_lines=('Суроолор барбы?', 'Сунуштар?'),
                       closing_email=None):
    """Returns a BytesIO of the finished .pptx (title/instructions slide,
    native-shape puzzle slide, closing slide)."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    _title_slide(prs, title, usage_lines)
    _puzzle_slide(prs, n, pairs)
    _closing_slide(prs, closing_lines, closing_email)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
