"""Builds an editable .pptx for a recall pack, mirroring the same grid
layout used by the PDF template (see recall_pack_print.html) so both
formats stay visually consistent, at slide proportions equivalent to the
reference pack (9906000 x 6858000 EMU)."""
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from .grid_layout import cell_rect_pct

SLIDE_W = Emu(9906000)
SLIDE_H = Emu(6858000)
MARGIN = Emu(180000)
HEADER_H = Emu(520000)

NAVY = RGBColor(0x0F, 0x3C, 0x58)
BLUE = RGBColor(0x0D, 0x6E, 0xFD)
GRAY = RGBColor(0x64, 0x74, 0x8B)
LIGHT_GRAY = RGBColor(0x94, 0xA3, 0xB8)
FONT = "Arial"

_BOLD_RE = re.compile(r"<strong>(.*?)</strong>", re.DOTALL)


def _cell_rect(number):
    """Converts the shared percentage-based grid geometry into an EMU
    rectangle within this slide's grid area (below the masthead)."""
    grid_left = MARGIN
    grid_top = MARGIN + HEADER_H
    grid_w = SLIDE_W - 2 * MARGIN
    grid_h = SLIDE_H - grid_top - MARGIN

    left_pct, top_pct, width_pct, height_pct = cell_rect_pct(number)
    left = Emu(int(grid_left + grid_w * left_pct / 100.0))
    top = Emu(int(grid_top + grid_h * top_pct / 100.0))
    width = Emu(int(grid_w * width_pct / 100.0))
    height = Emu(int(grid_h * height_pct / 100.0))
    return left, top, width, height


def _add_textbox(slide, left, top, width, height, lines, size=11, bold=False,
                  color=None, align=PP_ALIGN.LEFT, font=FONT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        _add_runs_from_html_lite(p, line, size, bold, color, font)
    return box


def _add_runs_from_html_lite(paragraph, text, size, bold, color, font):
    """Splits a <strong>...</strong>-tagged fragment into runs so bold
    spans render as real bold runs rather than literal tag text."""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            _add_run(paragraph, text[pos:m.start()], size, bold, color, font)
        _add_run(paragraph, m.group(1), size, True, color, font)
        pos = m.end()
    if pos < len(text):
        _add_run(paragraph, text[pos:], size, bold, color, font)


def _add_run(paragraph, text, size, bold, color, font):
    if not text:
        return
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    if color is not None:
        run.font.color.rgb = color


def _html_lines(html_text):
    return [seg for seg in html_text.split("<br>")]


def _add_cell_frame(slide, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = 0.03
    shape.fill.background()
    shape.line.color.rgb = NAVY
    shape.line.width = Pt(1.1)
    shape.shadow.inherit = False
    return shape


def _draw_reflection_diagram(slide, left, top, width, height, points, reflected, mirror_x, show_answer):
    rng_units = 5
    size = min(width, height) - Emu(200000)
    cell = Emu(int(size / (rng_units * 2)))
    ox = Emu(int(left + (width - cell * rng_units * 2) / 2))
    oy = Emu(int(top + (height - cell * rng_units * 2) / 2) + Emu(120000))

    def to_pt(x, y):
        return Emu(int(ox + (x + rng_units) * cell)), Emu(int(oy + (rng_units - y) * cell))

    for i in range(-rng_units, rng_units + 1):
        x0, y0 = to_pt(i, -rng_units)
        x1, y1 = to_pt(i, rng_units)
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x0, y0, x1, y1)
        conn.line.color.rgb = NAVY if i == 0 else LIGHT_GRAY
        conn.line.width = Pt(1.4 if i == 0 else 0.5)
        x0, y0 = to_pt(-rng_units, i)
        x1, y1 = to_pt(rng_units, i)
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x0, y0, x1, y1)
        conn.line.color.rgb = NAVY if i == 0 else LIGHT_GRAY
        conn.line.width = Pt(1.4 if i == 0 else 0.5)

    mx0, my0 = to_pt(mirror_x, -rng_units)
    mx1, my1 = to_pt(mirror_x, rng_units)
    mirror = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, mx0, my0, mx1, my1)
    mirror.line.color.rgb = RGBColor(0xDC, 0x26, 0x26)
    mirror.line.width = Pt(1.6)
    mirror.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    def add_triangle(pts, outline_rgb, fill_rgb, dashed):
        px, py = to_pt(*pts[0])
        fb = slide.shapes.build_freeform(start_x=px, start_y=py, scale=1.0)
        fb.add_line_segments([to_pt(*p) for p in pts[1:]] + [to_pt(*pts[0])], close=True)
        shape = fb.convert_to_shape()
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_rgb
        shape.line.color.rgb = outline_rgb
        shape.line.width = Pt(1.6)
        if dashed:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        shape.shadow.inherit = False
        return shape

    add_triangle(points, NAVY, RGBColor(0xDC, 0xE6, 0xEB), False)
    if show_answer:
        add_triangle(reflected, BLUE, RGBColor(0xDD, 0xE9, 0xFE), True)


def _add_question_cell(slide, q, show_answer):
    left, top, width, height = _cell_rect(q["number"])
    _add_cell_frame(slide, left, top, width, height)
    _add_textbox(slide, left + Emu(40000), top + Emu(20000), Emu(280000), Emu(280000),
                 q["circled"], size=14, bold=True, color=NAVY)
    _add_textbox(slide, left + width - Emu(360000), top + Emu(20000), Emu(320000), Emu(220000),
                 str(q["marks"]), size=10, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)
    _add_textbox(slide, left + Emu(260000), top + Emu(40000), width - Emu(320000), Emu(180000),
                 q["topic"], size=8, color=GRAY)

    is_geometry = "svg_points" in q
    prompt_h = height - Emu(700000) if is_geometry else height - Emu(280000)
    _add_textbox(slide, left + Emu(90000), top + Emu(220000), width - Emu(160000), prompt_h,
                 _html_lines(q["prompt"]), size=11)

    if is_geometry:
        _draw_reflection_diagram(
            slide, left + Emu(60000), top + Emu(900000), width - Emu(120000),
            height - Emu(980000), q["svg_points"], q["svg_reflected"],
            q["svg_mirror_x"], show_answer,
        )
    elif show_answer:
        _add_textbox(slide, left + Emu(90000), top + height - Emu(360000),
                     width - Emu(160000), Emu(320000), q["answer"], size=11,
                     bold=True, color=BLUE)


def _add_masthead(slide, pack, show_answer):
    _add_textbox(slide, MARGIN, Emu(60000), Emu(5000000), Emu(300000),
                 pack["series_title"], size=20, bold=True, color=NAVY)
    _add_textbox(slide, MARGIN, Emu(360000), Emu(5000000), Emu(200000),
                 pack["subject_line"], size=10, color=GRAY)
    _add_textbox(slide, SLIDE_W - MARGIN - Emu(1800000), Emu(60000), Emu(1800000), Emu(300000),
                 pack["pack_code"], size=18, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)
    flag = "Жооптор" if show_answer else "Суроолор"
    flag_color = BLUE if show_answer else NAVY
    _add_textbox(slide, SLIDE_W - MARGIN - Emu(1800000), Emu(360000), Emu(1800000), Emu(200000),
                 f"{flag}  ·  Макс. балл: {pack['total_marks']}", size=9,
                 bold=True, color=flag_color, align=PP_ALIGN.RIGHT)


def build_presentation(pack):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank_layout = prs.slide_layouts[6]

    for show_answer in (False, True):
        slide = prs.slides.add_slide(blank_layout)
        _add_masthead(slide, pack, show_answer)
        for q in pack["questions"]:
            _add_question_cell(slide, q, show_answer)

    closing = prs.slides.add_slide(blank_layout)
    bg = closing.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    bg.shadow.inherit = False
    _add_textbox(
        closing, Emu(600000), Emu(2600000), SLIDE_W - Emu(1200000), Emu(1600000),
        [
            "Сунуштарыңыз барбы?",
            "Ката тапсаңыз, бизге жазыңыз:",
            "support@mathkg.online",
        ],
        size=24, bold=False, color=RGBColor(0xFF, 0xFF, 0xFF),
        align=PP_ALIGN.CENTER,
    )
    return prs
